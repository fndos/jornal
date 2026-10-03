#!/usr/bin/env python3
"""Motor de calculo: biometrico -> formato de quincena.

No sabe nada de interfaz ni de linea de comandos. Lo usan por igual
`app.py` (interfaz grafica) y `llenar_quincena.py` (consola).
"""

import datetime as dt
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

TIPOS = ("DIARIO", "AFILIADO")

# Roles de persona: la celda del dia lleva siempre la letra, nunca un monto.
# G = guardias y Q = quemadores tienen horario nocturno y no siempre marcan
# completo. TRIUNFO es el afiliado enviado a esa sucursal, que ni siquiera marca.
ROLES = ("G", "Q", "TRIUNFO")

# Ausencias justificadas: son eventos, no roles, y NO se heredan entre quincenas.
LICENCIAS = ("PM", "IESS", "MSP", "VACACIONES", "PERMISO ENFERMEDAD")

CATEGORIAS_ACCION = {
    "SIN DATOS EN EL EXPORT", "CÉDULA DUPLICADA", "SIN CÉDULA", "CÉDULA IRREGULAR",
    "NO ESTÁ EN EL FORMATO", "SIN MARCACIONES", "MARCACIONES INCOMPLETAS",
    "REQUIERE CRITERIO", "RECONSTRUCCIÓN INSUFICIENTE", "RECONSTRUCCIÓN EN SÁBADO",
    "LICENCIA VIGENTE POR CONFIRMAR",
}

_ORDEN = [
    "SIN DATOS EN EL EXPORT", "CÉDULA DUPLICADA", "SIN CÉDULA", "CÉDULA IRREGULAR",
    "NO ESTÁ EN EL FORMATO", "SIN MARCACIONES", "LICENCIA VIGENTE POR CONFIRMAR",
    "CÉDULAS COMPLETADAS", "FÓRMULA CORREGIDA", "ROLES DEL PROPIO ARCHIVO",
    "ROL DE PERSONA", "MARCACIONES INCOMPLETAS",
    "REQUIERE CRITERIO", "RECONSTRUCCIÓN INSUFICIENTE", "RECONSTRUCCIÓN EN SÁBADO",
    "JORNADA PARCIAL", "MEDIA JORNADA", "JORNADA RECONSTRUIDA",
    "PRESENCIA SIN HORAS", "POSIBLE HORA EXTRA",
]


# --------------------------------------------------------------------------- #
# cedulas
# --------------------------------------------------------------------------- #
def normalizar_cedula(valor):
    """La cedula ecuatoriana tiene 10 digitos. Excel se come el 0 inicial de las
    provincias 01 a 09 al guardarlas como numero, asi que hay que devolverselo.

    Solo se rellena el caso inequivoco de 9 digitos. Cualquier otra longitud se
    deja tal cual y se reporta, porque rellenarla seria inventar.
    """
    if valor is None:
        return ""
    limpia = "".join(c for c in str(valor).strip() if c.isdigit())
    if len(limpia) == 9:
        return "0" + limpia
    return limpia


def cedula_valida(cedula):
    """Verifica provincia y digito verificador (modulo 10)."""
    if len(cedula) != 10 or not cedula.isdigit():
        return False
    provincia = int(cedula[:2])
    if not (1 <= provincia <= 24 or provincia == 30):
        return False
    if int(cedula[2]) > 5:
        return False
    total = 0
    for i, d in enumerate(cedula[:9]):
        n = int(d) * (2 if i % 2 == 0 else 1)
        total += n - 9 if n > 9 else n
    return (10 - total % 10) % 10 == int(cedula[9])


# --------------------------------------------------------------------------- #
# utilidades
# --------------------------------------------------------------------------- #
def a_segundos(t):
    p = [int(x) for x in str(t).strip().split(":")]
    while len(p) < 3:
        p.append(0)
    return p[0] * 3600 + p[1] * 60 + p[2]


def hhmm(s):
    return f"{int(s) // 3600:02d}:{int(s) % 3600 // 60:02d}"


def a_fecha(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return dt.date.fromisoformat(str(v).strip()[:10])


def redondear(h, paso):
    return round(h / paso) * paso


def fmt_horas(h):
    return f"{h:g}"


def dinero(valor):
    """Redondea a centavos, al alza en el empate, y devuelve un numero.

    Se escribe el importe ya resuelto y no la formula: openpyxl no puede
    guardar el resultado junto a la formula, asi que un visor que no recalcula
    muestra 0 en esa celda. Ademas es como lo escriben a mano: 10.1, no 10.095.
    """
    d = Decimal(str(valor)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return int(d) if d == d.to_integral_value() else float(d)


# --------------------------------------------------------------------------- #
# lectura
# --------------------------------------------------------------------------- #
def _sin_tildes(v):
    t = str(v or "").strip().lower()
    for a, b in zip("áéíóúü", "aeiouu"):
        t = t.replace(a, b)
    return t


def columnas_biometrico(ws):
    """Ubica las columnas del biometrico por el nombre de su cabecera.

    El reloj no siempre exporta las mismas columnas: una version traia Genero
    y otra no, y eso corre todo de lugar. Buscar por nombre evita depender de
    la posicion. Si la cabecera no se entiende, se deduce mirando los datos.
    """
    cab = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    idx = {}
    for j, bruto in enumerate(cab):
        n = _sin_tildes(bruto)
        if not n:
            continue
        if "fecha" in n:
            idx.setdefault("fecha", j)
        elif n.startswith("hora"):
            idx.setdefault("hora", j)
        elif "marcacion" in n or "numero" in n:
            continue                       # el contador de marcaciones no se usa
        elif n.startswith("id") or "identific" in n or "cedula" in n:
            idx.setdefault("id", j)
        elif "departamento" in n or "area" in n:
            idx.setdefault("departamento", j)

    if "fecha" not in idx or "hora" not in idx:
        idx.update(_columnas_por_contenido(ws, idx))

    faltan = [k for k in ("id", "fecha", "hora") if k not in idx]
    if faltan:
        nombres = ", ".join(str(c) for c in cab if c) or "(sin cabecera)"
        raise ValueError(
            "No encuentro la columna de " + " ni de ".join(faltan) +
            f" en la hoja del biometrico. Sus columnas son: {nombres}.")
    return idx


def _columnas_por_contenido(ws, ya):
    """Cuando la cabecera no sirve, se mira la primera fila con datos: la fecha
    es lo que se puede leer como fecha y la hora lo que tiene formato hh:mm:ss."""
    import re
    encontrado = {}
    for fila in ws.iter_rows(min_row=2, max_row=40, values_only=True):
        if not fila or all(v is None for v in fila):
            continue
        for j, v in enumerate(fila):
            if v is None:
                continue
            if "fecha" not in ya and "fecha" not in encontrado:
                try:
                    a_fecha(v)
                    encontrado["fecha"] = j
                    continue
                except Exception:
                    pass
            if "hora" not in ya and "hora" not in encontrado:
                if re.fullmatch(r"\s*\d{1,2}:\d{2}(:\d{2})?(\s*,\s*\d{1,2}:\d{2}(:\d{2})?)*\s*",
                                str(v)):
                    encontrado["hora"] = j
        if {"fecha", "hora"} <= (set(ya) | set(encontrado)):
            break
    encontrado.setdefault("id", 0)
    return encontrado


def leer_biometrico(ws):
    """-> ({cedula: {fecha: [segundos,...]}}, {cedula: departamento})"""
    col = columnas_biometrico(ws)
    marcas = defaultdict(dict)
    depto = {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila or len(fila) <= col["hora"] or fila[col["id"]] is None:
            continue
        ced = normalizar_cedula(fila[col["id"]])
        try:
            fecha = a_fecha(fila[col["fecha"]])
        except Exception:
            continue
        segs = sorted(a_segundos(x) for x in str(fila[col["hora"]] or "").split(",") if x.strip())
        marcas[ced][fecha] = segs
        if "departamento" in col:
            depto[ced] = str(fila[col["departamento"]] or "").strip()
    return marcas, depto


MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


def periodo_de(fecha):
    """A que quincena pertenece una fecha: (primer dia, ultimo dia)."""
    if fecha.day <= 15:
        return fecha.replace(day=1), fecha.replace(day=15)
    if fecha.month == 12:
        siguiente = fecha.replace(year=fecha.year + 1, month=1, day=1)
    else:
        siguiente = fecha.replace(month=fecha.month + 1, day=1)
    return fecha.replace(day=16), siguiente - dt.timedelta(days=1)


def titulo_periodo(inicio, fin):
    return f"{inicio.day} AL {fin.day} DE {MESES[inicio.month - 1]} DEL {inicio.year}"


def identificar_hojas(archivo):
    """-> (hoja del biometrico, hoja del formato o None).

    Se decide por contenido, no por posicion: el mismo libro llega unas veces
    con el biometrico primero y otras con el formato primero.
    """
    wb = openpyxl.load_workbook(archivo, data_only=False)
    try:
        bio = fmt = None
        for nombre in wb.sheetnames:
            ws = wb[nombre]
            if fmt is None:
                try:
                    filas, _ = leer_formato(ws)
                    if len(filas) >= 10:
                        fmt = nombre
                        continue
                except Exception:
                    pass
            if bio is None:
                try:
                    columnas_biometrico(ws)
                    bio = nombre
                except Exception:
                    pass
        if bio is None:
            raise ValueError(
                "Ninguna hoja de este libro parece el registro biometrico. "
                f"Sus hojas son: {', '.join(wb.sheetnames)}.")
        return bio, fmt
    finally:
        wb.close()


def hoja_de_formato(archivo):
    """Nombre de la hoja del formato dentro del libro, si la trae."""
    try:
        return identificar_hojas(archivo)[1]
    except Exception:
        return None


def primera_fecha(archivo, hoja=None):
    """Lee solo la columna de fechas del biometrico para saber de que quincena
    es, sin cargar el libro entero."""
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    try:
        ws = wb[hoja or identificar_hojas(archivo)[0]]
        j = columnas_biometrico(ws)["fecha"]
        for fila in ws.iter_rows(min_row=2, min_col=j + 1, max_col=j + 1, values_only=True):
            try:
                return a_fecha(fila[0])
            except Exception:
                continue
    finally:
        wb.close()
    return None


def nombre_sugerido(fecha):
    """Nombre de archivo para la quincena a la que pertenece `fecha`.

    Se antepone el año y el mes para que la carpeta de resultados quede
    ordenada cronológicamente sola.
    """
    inicio, fin = periodo_de(fecha)
    return f"Formato_{inicio.year}_{inicio.month:02d}_{inicio.day:02d}-{fin.day:02d}.xlsx"


def roles_en_formato(ws, filas):
    """Roles y licencias que el formato ya trae escritos. Quien prepara el
    archivo suele dejarlos puestos, y entonces no hace falta ninguna tabla
    guardada: la fuente mas fiable es el propio archivo."""
    cols = [5 + 2 * k for k in range(16)]
    out = {}
    for f in filas:
        if not f["cedula"]:
            continue
        vistos = {}
        for c in cols:
            v = str(ws.cell(f["fila"], c).value or "").strip().upper()
            if v in ROLES + LICENCIAS:
                vistos[v] = vistos.get(v, 0) + 1
        if vistos:
            cod, n = max(vistos.items(), key=lambda kv: kv[1])
            if n >= 3:
                out[f["cedula"]] = cod
    return out


def leer_formato(ws):
    """Detecta las filas de personal y las 16 columnas de dia. No modifica nada."""
    filas = []
    for i in range(1, ws.max_row + 1):
        tipo = ws.cell(i, 4).value
        if isinstance(tipo, str) and tipo.strip().upper() in TIPOS:
            ced = ws.cell(i, 2).value
            nombre = ws.cell(i, 3).value
            filas.append({
                "fila": i,
                "cedula": normalizar_cedula(ced) or None,
                "cedula_original": str(ced).strip() if ced is not None else None,
                "nombre": str(nombre).strip() if nombre else "",
                "tipo": tipo.strip().upper(),
            })
    if not filas:
        raise ValueError("La hoja no tiene filas con DIARIO o AFILIADO en la columna D.")

    cabecera = next(
        (i for i in range(1, ws.max_row + 1) if ws.cell(i, 2).value == "Cedula/ID"), None
    )
    if cabecera is None:
        raise ValueError("No encontre la cabecera 'Cedula/ID' en la columna B.")
    return filas, cabecera


# --------------------------------------------------------------------------- #
# reconstruccion de la jornada
# --------------------------------------------------------------------------- #
def clasificar(segs, cfg):
    """Agrupa las marcaciones en ENTRADA / ALM_OUT / MEDIODIA / ALM_IN / SALIDA."""
    lim = {k: a_segundos(v) for k, v in cfg["ventanas"].items()}
    dedupe = cfg["minutos_marca_duplicada"] * 60

    limpias = []
    for s in segs:
        if limpias and s - limpias[-1] <= dedupe:
            continue
        limpias.append(s)

    b = {}
    for i, s in enumerate(limpias):
        ultima = i == len(limpias) - 1
        if s < lim["fin_entrada"]:
            b.setdefault("ENTRADA", []).append(s)
        elif s < lim["fin_alm_out"]:
            b.setdefault("ALM_OUT", []).append(s)
        elif s < lim["fin_mediodia"]:
            b.setdefault("MEDIODIA" if ultima else "ALM_OUT", []).append(s)
        elif s < lim["fin_alm_in"]:
            b.setdefault("ALM_IN", []).append(s)
        else:
            b.setdefault("SALIDA", []).append(s)

    # secuencia real del dia: se entra una vez, se sale a almorzar al final de la
    # manana, se vuelve al principio de la tarde y se sale al final del dia.
    return {
        "ENTRADA": min(b["ENTRADA"]) if "ENTRADA" in b else None,
        "ALM_OUT": max(b["ALM_OUT"]) if "ALM_OUT" in b else None,
        "MEDIODIA": max(b["MEDIODIA"]) if "MEDIODIA" in b else None,
        "ALM_IN": min(b["ALM_IN"]) if "ALM_IN" in b else None,
        "SALIDA": max(b["SALIDA"]) if "SALIDA" in b else None,
    }, limpias


def perfil_persona(dias, cfg):
    """Horario habitual, tomando siempre el extremo que acorta la jornada."""
    acum = defaultdict(list)
    for segs in dias.values():
        b, _ = clasificar(segs, cfg)
        if all(b[k] is not None for k in ("ENTRADA", "ALM_OUT", "ALM_IN", "SALIDA")):
            for k in ("ENTRADA", "ALM_OUT", "ALM_IN", "SALIDA"):
                acum[k].append(b[k])
    if not acum:
        return None
    return {
        "ENTRADA": max(acum["ENTRADA"]),
        "ALM_OUT": min(acum["ALM_OUT"]),
        "ALM_IN": max(acum["ALM_IN"]),
        "SALIDA": min(acum["SALIDA"]),
    }


def calcular_dia(segs, cfg, perfil=None):
    """-> (horas | None, calidad, detalle)"""
    if not segs:
        return None, "sin_marcas", "sin marcaciones"

    b, limpias = clasificar(segs, cfg)
    marcas = ",".join(hhmm(s) for s in limpias)

    if b["MEDIODIA"] and b["ENTRADA"] and not b["ALM_IN"] and not b["SALIDA"]:
        return (b["MEDIODIA"] - b["ENTRADA"]) / 3600, "media_jornada", f"mañana ({marcas})"

    faltan = [k for k in ("ENTRADA", "ALM_OUT", "ALM_IN", "SALIDA") if b[k] is None]

    if not faltan:
        neto = ((b["ALM_OUT"] - b["ENTRADA"]) + (b["SALIDA"] - b["ALM_IN"])) / 3600
        return neto, "exacta", f"4 marcaciones ({marcas})"

    if len(faltan) == 1 and perfil and cfg["reconstruir_marca_faltante"]:
        k = faltan[0]
        est = dict(b)
        est[k] = perfil[k]
        neto = ((est["ALM_OUT"] - est["ENTRADA"]) + (est["SALIDA"] - est["ALM_IN"])) / 3600
        if neto <= 0:
            return None, "incalculable", f"reconstrucción inconsistente ({marcas})"
        return neto, "reconstruida", f"falta {k}, estimada en {hhmm(perfil[k])} ({marcas})"

    return None, "incalculable", f"faltan {'+'.join(faltan)} ({marcas})"


def valor_dia(horas, tipo, fecha, cfg):
    """-> (valor | None, nota)"""
    dow = fecha.weekday()
    h = redondear(horas, cfg["redondeo_horas"])

    if dow == 6:
        return None, "domingo: tarifa no deducible del histórico"

    if dow == 5:
        if not cfg["escribir_sabados"]:
            return None, "sábado omitido por configuración"
        if abs(h - cfg["sabado_estandar_horas"]) < 1e-9:
            return cfg["sabado_estandar_valor"], "sábado estándar"
        return dinero(h * cfg["tarifa_hora_finde"]), "sábado por horas"

    if tipo == "AFILIADO":
        return cfg["marca_afiliado_presente"], "afiliado presente"

    if h >= cfg["horas_jornada_completa"]:
        return cfg["jornal_dia_completo"], "jornada completa"

    tar = cfg["tarifa_hora_semana"]
    bruto = h * tar
    if abs(h - int(h)) >= 1e-9:          # las medias horas llevan el ajuste
        bruto += cfg["ajuste_media_hora"]
    return dinero(bruto), "jornada parcial"


# --------------------------------------------------------------------------- #
# proceso completo
# --------------------------------------------------------------------------- #
def ruta_revision(salida):
    """El informe de revision va aparte, para que el archivo que recibe la
    jefatura conserve una sola hoja."""
    salida = Path(salida)
    nombre = salida.stem
    if nombre.startswith("Formato_"):
        nombre = "Revision_" + nombre[len("Formato_"):]
    else:
        nombre = f"{nombre}_Revision"
    return salida.with_name(f"{nombre}.xlsx")


def procesar(archivo, salida, cfg, plantilla=None, colaboradores=None,
             hoja_bio=None, hoja_formato=None, salida_revision=None, progreso=None):
    """Llena el formato de la jefatura y devuelve un resumen.

    archivo   : libro del biometrico (la hoja de marcaciones).
    plantilla : libro con el formato de la jefatura. Si se omite, se usa la
                segunda hoja de `archivo` (el caso de un libro con las dos).
    salida    : archivo nuevo con el formato lleno. Conserva una sola hoja.
    El informe de incidencias se guarda en un archivo aparte.

    Ninguno de los archivos de entrada se modifica.
    """
    def avisar(txt):
        if progreso:
            progreso(txt)

    colaboradores = dict(colaboradores or {})
    avisar("Abriendo el biométrico...")
    if hoja_bio is None:
        hoja_bio = identificar_hojas(archivo)[0]
    wb_bio = openpyxl.load_workbook(archivo, data_only=True)
    ws_bio = wb_bio[hoja_bio]

    if plantilla:
        avisar("Abriendo la plantilla...")
        wb = openpyxl.load_workbook(plantilla, data_only=False)
        ws = wb[hoja_formato or wb.sheetnames[0]]
    else:
        # un solo libro con biometrico en la hoja 1 y formato en la hoja 2
        wb = openpyxl.load_workbook(archivo, data_only=False)
        if len(wb.sheetnames) < 2:
            raise ValueError(
                "El libro trae una sola hoja. Indica tambien el archivo con el "
                "formato de la jefatura.")
        ws = wb[hoja_formato or wb.sheetnames[1]]

    avisar("Leyendo marcaciones...")
    marcas, depto = leer_biometrico(ws_bio)
    filas, cabecera = leer_formato(ws)
    fechas_export = sorted({f for d in marcas.values() for f in d})
    if not fechas_export:
        raise ValueError("El biometrico no tiene ninguna marcacion legible.")

    # El periodo sale del biometrico, no de la plantilla: asi la misma plantilla
    # sirve todas las quincenas. Las demas fechas de la cabecera son formulas
    # (=+E2+1), de modo que se acomodan solas al cambiar la primera.
    inicio, fin_quincena = periodo_de(fechas_export[0])
    if cfg.get("ajustar_fechas_del_periodo", True):
        ws.cell(cabecera, 5).value = dt.datetime(inicio.year, inicio.month, inicio.day)
        titulo = ws.cell(cabecera - 1, 5)
        if cabecera > 1 and isinstance(titulo.value, str):
            titulo.value = titulo_periodo(inicio, fin_quincena)
    else:
        base = ws.cell(cabecera, 5).value
        if isinstance(base, (dt.date, dt.datetime)):
            inicio, fin_quincena = periodo_de(a_fecha(base))

    columnas = [(5 + 2 * k, inicio + dt.timedelta(days=k)) for k in range(16)]

    del_formato = roles_en_formato(ws, filas)
    for ced, cod in del_formato.items():
        ficha = dict(colaboradores.get(ced, {}))
        ficha["rol"] = cod if cod in ROLES else ""
        ficha["licencia"] = cod if cod in LICENCIAS else ""
        colaboradores[ced] = ficha
    todas = {(c, f): s for c, d in marcas.items() for f, s in d.items()}
    perfil_global = perfil_persona(todas, cfg)

    por_cedula = defaultdict(list)
    for f in filas:
        if f["cedula"]:
            por_cedula[f["cedula"]].append(f)

    incidencias = []
    escritas = Counter()
    respetar = cfg.get("respetar_lo_ya_escrito", True)
    ya_escritas = {(f["fila"], c) for f in filas for c, _ in columnas
                   if ws.cell(f["fila"], c).value is not None} if respetar else set()

    def escribir(fila, col, valor, concepto):
        """Nunca encima de algo que ya estaba: esas celdas las puso una persona."""
        if (fila, col) in ya_escritas:
            escritas["respetado lo que ya estaba"] += 1
            return False
        ws.cell(fila, col).value = valor
        escritas[concepto] += 1
        return True

    def anotar(cat, f, fecha, detalle, dato=""):
        incidencias.append({
            "categoria": cat,
            "fila": f["fila"] if f else "",
            "nombre": f["nombre"] if f else "",
            "tipo": f["tipo"] if f else "",
            "fecha": fecha.isoformat() if fecha else "",
            "detalle": detalle,
            "dato": dato,
        })

    if del_formato:
        anotar("ROLES DEL PROPIO ARCHIVO", None, None,
               f"El formato ya traía escritos los roles o licencias de "
               f"{len(del_formato)} personas; se usaron esos en vez de la tabla guardada.")

    # --- problemas estructurales ---
    corregir = cfg.get("corregir_cedulas_en_el_formato", True)
    rellenadas = 0
    for f in filas:
        if not f["cedula"]:
            anotar("SIN CÉDULA", f, None, "La fila no tiene cédula: imposible cruzar")
            continue
        if f["cedula"] != f.get("cedula_original"):
            rellenadas += 1
            if corregir:
                ws.cell(f["fila"], 2).value = f["cedula"]
        if not cedula_valida(f["cedula"]):
            anotar("CÉDULA IRREGULAR", f, None,
                   f"'{f['cedula']}' no pasa la validación del dígito verificador "
                   f"ecuatoriano. Verificar contra el documento.")
    if rellenadas:
        anotar("CÉDULAS COMPLETADAS", None, None,
               f"{rellenadas} cédulas venían sin el 0 inicial (Excel lo quita al "
               f"guardarlas como número). " +
               ("Se escribieron completas en el formato."
                if corregir else
                "En el formato quedaron como estaban; se pueden corregir activando "
                "«Completar las cédulas en el formato»."))
    for ced, fs in por_cedula.items():
        if len(fs) > 1:
            for f in fs:
                anotar("CÉDULA DUPLICADA", f, None,
                       f"La cédula {ced} aparece en {len(fs)} filas; se omiten todas")

    dias_periodo = (fin_quincena - inicio).days + 1
    for k in range(dias_periodo):
        d = inicio + dt.timedelta(days=k)
        if d not in fechas_export and d.weekday() != 6:
            anotar("SIN DATOS EN EL EXPORT", None, d, "Ningún registro biométrico ese día")

    for ced in sorted(set(marcas) - set(por_cedula)):
        anotar("NO ESTÁ EN EL FORMATO", None, None,
               f"Marca biométrico pero no existe fila; {len(marcas[ced])} días marcados",
               f"cédula {ced} / depto {depto.get(ced, '')}")

    # --- llenado ---
    avisar(f"Procesando {len(filas)} personas...")
    for f in filas:
        ced = f["cedula"]
        if not ced or len(por_cedula[ced]) > 1:
            continue

        ficha = colaboradores.get(ced, {})
        dias = marcas.get(ced, {})
        perfil = perfil_persona(dias, cfg) or perfil_global

        # Guardias, quemadores y El Triunfo: la celda lleva la letra todos los
        # dias, marquen o no. Su jornada es nocturna o pactada y el biometrico
        # no la describe, asi que no se calcula ningun monto.
        rol = (ficha.get("rol") or "").strip().upper()
        if rol in ROLES:
            for col, fecha in columnas:
                if not (inicio <= fecha <= fin_quincena):
                    continue
                if fecha.weekday() == 6 and fecha not in dias:
                    continue
                escribir(f["fila"], col, rol, f"rol {rol}")
            anotar("ROL DE PERSONA", f, None,
                   f"Rol '{rol}' del registro de colaboradores; se escribe toda la "
                   f"quincena (marcó {len(dias)} días del export)")
            continue

        # Un permiso medico o unas vacaciones son un evento, no un estado fijo.
        lic = (ficha.get("licencia") or "").strip().upper()
        if lic in LICENCIAS:
            anotar("LICENCIA VIGENTE POR CONFIRMAR", f, None,
                   f"El registro lo marca como '{lic}'. Confirmar vigencia y escribirlo "
                   f"a mano. Marcaciones en este período: {len(dias)}")
            if not dias:
                continue

        if not dias:
            anotar("SIN MARCACIONES", f, None,
                   "No aparece en el biométrico en todo el período; fila en blanco")
            continue

        for col, fecha in columnas:
            if fecha > fin_quincena or fecha not in fechas_export:
                continue

            celda = ws.cell(f["fila"], col)
            segs = dias.get(fecha)

            if segs is None:
                if fecha.weekday() == 6 or not cfg["escribir_ausencias"]:
                    continue
                escribir(f["fila"], col, cfg["marca_ausencia"], "ausencia")
                continue

            horas, calidad, detalle = calcular_dia(segs, cfg, perfil)
            crudo = ",".join(hhmm(s) for s in segs)

            if horas is None and f["tipo"] == "AFILIADO" and fecha.weekday() < 5:
                escribir(f["fila"], col, cfg["marca_afiliado_presente"], "afiliado presente")
                anotar("PRESENCIA SIN HORAS", f, fecha,
                       "Marcaciones incompletas, pero consta presente: " + detalle, crudo)
                continue

            if horas is None:
                anotar("MARCACIONES INCOMPLETAS", f, fecha, detalle, crudo)
                continue

            if calidad == "reconstruida" and f["tipo"] == "DIARIO":
                if fecha.weekday() < 5 and redondear(horas, cfg["redondeo_horas"]) < cfg["horas_jornada_completa"]:
                    anotar("RECONSTRUCCIÓN INSUFICIENTE", f, fecha,
                           f"{detalle}; da {horas:.2f} h, no alcanza la jornada completa", crudo)
                    continue
                if fecha.weekday() == 5:
                    anotar("RECONSTRUCCIÓN EN SÁBADO", f, fecha, detalle, crudo)
                    continue

            valor, nota = valor_dia(horas, f["tipo"], fecha, cfg)
            if valor is None:
                anotar("REQUIERE CRITERIO", f, fecha, nota, f"{horas:.2f} h / {crudo}")
                continue

            if not escribir(f["fila"], col, valor, nota):
                continue

            if calidad == "reconstruida":
                anotar("JORNADA RECONSTRUIDA", f, fecha, detalle, f"neto {horas:.2f} h")
            elif calidad == "media_jornada":
                anotar("MEDIA JORNADA", f, fecha, f"{detalle}; {horas:.2f} h netas", crudo)
            if nota == "jornada parcial":
                anotar("JORNADA PARCIAL", f, fecha,
                       f"{horas:.2f} h netas → se pagó por horas; verificar contra la bitácora", crudo)

            corte = a_segundos(cfg["umbral_candidato_extras_salida"])
            if segs[-1] >= corte or horas >= cfg["umbral_candidato_extras_horas"]:
                anotar("POSIBLE HORA EXTRA", f, fecha,
                       f"salida {hhmm(segs[-1])}, {horas:.2f} h netas; confirmar con la bitácora", crudo)

    # --- subtotal ---
    if cfg["corregir_formula_subtotal"] and filas:
        pri, ult = filas[0]["fila"], filas[-1]["fila"]
        celda = ws.cell(ult + 1, 38)
        esperado = f"=SUM(AL{pri}:AL{ult})"
        if isinstance(celda.value, str) and celda.value.upper().startswith("=SUM"):
            if celda.value.replace(" ", "").upper() != esperado.upper():
                anotar("FÓRMULA CORREGIDA", None, None,
                       f"Subtotal AL{ult + 1} decía {celda.value}; se corrigió a {esperado}")
                celda.value = esperado

    avisar("Guardando el formato...")
    # El archivo que recibe la jefatura lleva una sola hoja. Cuando el formato
    # venia dentro del libro del biometrico, hay que soltar las demas.
    for nombre in list(wb.sheetnames):
        if nombre != ws.title:
            del wb[nombre]
    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)

    avisar("Guardando el informe de revisión...")
    salida_revision = Path(salida_revision or ruta_revision(salida))
    _libro_revision(incidencias).save(salida_revision)

    return {
        "salida": str(salida),
        "salida_revision": str(salida_revision),
        "origen_formato": "del archivo del biométrico" if plantilla == str(archivo)
                          else "de la plantilla guardada",
        "inicio": inicio,
        "fin": fin_quincena,
        "cobertura": (fechas_export[0], fechas_export[-1], len(fechas_export)),
        "personas": len(filas),
        "con_marcaciones": sum(1 for f in filas if f["cedula"] in marcas),
        "escritas": escritas,
        "incidencias": incidencias,
        "resumen_incidencias": Counter(i["categoria"] for i in incidencias),
    }


def _libro_revision(incidencias):
    """Libro aparte con todo lo que quedo por decidir o por verificar."""
    wb = openpyxl.Workbook()
    rev = wb.active
    rev.title = "REVISION"
    rev.append(["Accion", "Categoria", "Fila", "Personal", "Tipo",
                "Fecha", "Detalle", "Marcaciones"])
    for c in rev[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="FF404040")
        c.alignment = Alignment(horizontal="center")

    rojo = PatternFill("solid", fgColor="FFFFC7CE")
    orden = {c: i for i, c in enumerate(_ORDEN)}
    for it in sorted(incidencias,
                     key=lambda x: (orden.get(x["categoria"], 99), str(x["fecha"]), x["fila"] or 0)):
        accion = "LLENAR A MANO" if it["categoria"] in CATEGORIAS_ACCION else "informativo"
        rev.append([accion, it["categoria"], it["fila"], it["nombre"], it["tipo"],
                    it["fecha"], it["detalle"], it["dato"]])
        if accion == "LLENAR A MANO":
            rev.cell(rev.max_row, 1).fill = rojo

    for col, ancho in zip("ABCDEFGH", (16, 30, 7, 36, 10, 12, 62, 40)):
        rev.column_dimensions[col].width = ancho
    rev.freeze_panes = "A2"
    rev.auto_filter.ref = f"A1:H{rev.max_row}"
    return wb
